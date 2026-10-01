from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from math import fsum, isfinite
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_NAME = "Central_Superstore.xlsx"
SHEET_NAME = "Central_Region"
EXPECTED_COLUMNS = [
    "Row ID",
    "Order ID",
    "Order Date",
    "Ship Date",
    "Ship Mode",
    "Customer ID",
    "Customer Name",
    "Segment",
    "Country",
    "City",
    "State",
    "Postal Code",
    "Region",
    "Product ID",
    "Category",
    "Sub-Category",
    "Product Name",
    "Sales",
    "Quantity",
    "Discount",
    "Profit",
]

DESCRIPTIONS = {
    "Row ID": "Source-row identifier.",
    "Order ID": "Identifier shared by product lines in the same order.",
    "Order Date": "Date the customer placed the order.",
    "Ship Date": "Date the order was shipped.",
    "Ship Mode": "Shipping service selected for the order.",
    "Customer ID": "Identifier for the customer.",
    "Customer Name": "Customer's displayed name.",
    "Segment": "Customer business segment.",
    "Country": "Country associated with the order location.",
    "City": "City associated with the order location.",
    "State": "State associated with the order location.",
    "Postal Code": "Postal code associated with the order location.",
    "Region": "Sales region associated with the order.",
    "Product ID": "Identifier for the product.",
    "Category": "High-level product category.",
    "Sub-Category": "Product sub-category.",
    "Product Name": "Product's displayed name.",
    "Sales": "Sales amount recorded for this order line.",
    "Quantity": "Number of units on this order line.",
    "Discount": "Discount rate recorded for this order line.",
    "Profit": "Profit or loss recorded for this order line.",
}


def is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def display(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.12g}"
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def markdown_number(value: float | int | None) -> str:
    if value is None:
        return "n/a"
    return f"{value:,.2f}"


def infer_type(values: list[Any]) -> str:
    observed = [value for value in values if not is_missing(value)]
    if not observed:
        return "No non-missing values"
    kinds = set()
    for value in observed:
        if isinstance(value, datetime):
            kinds.add("date/datetime")
        elif isinstance(value, date):
            kinds.add("date")
        elif isinstance(value, bool):
            kinds.add("boolean")
        elif isinstance(value, int):
            kinds.add("integer")
        elif isinstance(value, float):
            kinds.add("number (decimal)")
        elif isinstance(value, str):
            kinds.add("text")
        else:
            kinds.add(type(value).__name__)
    return ", ".join(sorted(kinds))


def format_values(values: set[Any], limit: int = 12) -> str:
    ordered = sorted(values, key=lambda value: display(value).casefold())
    rendered = [display(value) for value in ordered[:limit]]
    if len(ordered) > limit:
        rendered.append(f"... (+{len(ordered) - limit} more)")
    return ", ".join(rendered)


def grouped_values(rows: list[tuple[Any, ...]], key_index: int, value_index: int) -> dict[Any, set[Any]]:
    groups: dict[Any, set[Any]] = defaultdict(set)
    for row in rows:
        key, value = row[key_index], row[value_index]
        if not is_missing(key) and not is_missing(value):
            groups[key].add(value)
    return groups


def inconsistent_keys(groups: dict[Any, set[Any]]) -> dict[Any, set[Any]]:
    return {key: values for key, values in groups.items() if len(values) > 1}


def repeated_key_summary(rows: list[tuple[Any, ...]], key_index: int, other_index: int | None = None) -> tuple[int, int, int]:
    by_key: dict[Any, set[Any]] = defaultdict(set)
    counts: Counter[Any] = Counter()
    for row in rows:
        key = row[key_index]
        if is_missing(key):
            continue
        counts[key] += 1
        if other_index is not None:
            by_key[key].add(row[other_index])
    multiple_rows = sum(count > 1 for count in counts.values())
    multiple_other = sum(len(values) > 1 for values in by_key.values()) if other_index is not None else 0
    return len(counts), multiple_rows, multiple_other


def find_workbook() -> Path:
    candidates = [PROJECT_ROOT / "data" / WORKBOOK_NAME, PROJECT_ROOT / WORKBOOK_NAME]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    searched = ", ".join(str(path) for path in candidates)
    raise FileNotFoundError(f"Could not find {WORKBOOK_NAME}; searched: {searched}")


def main() -> None:
    workbook_path = find_workbook()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Expected sheet {SHEET_NAME!r}; found {workbook.sheetnames!r}")

    worksheet = workbook[SHEET_NAME]
    iterator = worksheet.iter_rows(values_only=True)
    headers_raw = next(iterator, None)
    if not headers_raw:
        raise ValueError(f"Worksheet {SHEET_NAME!r} is empty")
    headers = [str(value).strip() if value is not None else "" for value in headers_raw]
    if headers != EXPECTED_COLUMNS:
        raise ValueError(f"Unexpected columns. Expected {EXPECTED_COLUMNS!r}; found {headers!r}")

    rows = [tuple(row[:len(headers)]) for row in iterator if any(not is_missing(value) for value in row)]
    workbook.close()
    column_index = {name: index for index, name in enumerate(headers)}
    columns = {name: [row[index] for row in rows] for name, index in column_index.items()}
    missing_counts = {name: sum(is_missing(value) for value in values) for name, values in columns.items()}
    unique_counts = {name: len({value for value in values if not is_missing(value)}) for name, values in columns.items()}
    duplicate_rows = len(rows) - len(set(rows))

    order_index = column_index["Order ID"]
    customer_index = column_index["Customer ID"]
    product_index = column_index["Product ID"]
    order_count, orders_with_multiple_rows, _ = repeated_key_summary(rows, order_index)
    customer_count, customers_with_multiple_rows, customers_with_multiple_orders = repeated_key_summary(
        rows, customer_index, order_index
    )
    product_count, products_with_multiple_rows, products_with_multiple_orders = repeated_key_summary(
        rows, product_index, order_index
    )

    customer_name_conflicts = inconsistent_keys(grouped_values(rows, customer_index, column_index["Customer Name"]))
    customer_segment_conflicts = inconsistent_keys(grouped_values(rows, customer_index, column_index["Segment"]))
    product_name_conflicts = inconsistent_keys(grouped_values(rows, product_index, column_index["Product Name"]))
    product_category_conflicts = inconsistent_keys(grouped_values(rows, product_index, column_index["Category"]))
    product_subcategory_conflicts = inconsistent_keys(grouped_values(rows, product_index, column_index["Sub-Category"]))

    orders = defaultdict(lambda: {"customers": set(), "order_dates": set(), "ship_dates": set(), "ship_modes": set()})
    for row in rows:
        order = row[order_index]
        if not is_missing(order):
            orders[order]["customers"].add(row[column_index["Customer ID"]])
            orders[order]["order_dates"].add(row[column_index["Order Date"]])
            orders[order]["ship_dates"].add(row[column_index["Ship Date"]])
            orders[order]["ship_modes"].add(row[column_index["Ship Mode"]])
    inconsistent_orders = {
        key: values
        for key, values in orders.items()
        if any(len(value_set) > 1 for value_set in values.values())
    }

    order_dates = [value for value in columns["Order Date"] if isinstance(value, (date, datetime))]
    ship_dates = [value for value in columns["Ship Date"] if isinstance(value, (date, datetime))]
    invalid_order_dates = len(rows) - len(order_dates)
    invalid_ship_dates = len(rows) - len(ship_dates)
    ship_before_order = [
        (row[order_index], row[column_index["Order Date"]], row[column_index["Ship Date"]])
        for row in rows
        if isinstance(row[column_index["Order Date"]], (date, datetime))
        and isinstance(row[column_index["Ship Date"]], (date, datetime))
        and row[column_index["Ship Date"]] < row[column_index["Order Date"]]
    ]

    numeric_stats: dict[str, dict[str, float | int | None]] = {}
    for name in ("Sales", "Quantity", "Discount", "Profit"):
        numbers = [value for value in columns[name] if isinstance(value, (int, float)) and not isinstance(value, bool)]
        numeric_stats[name] = {
            "count": len(numbers),
            "min": min(numbers) if numbers else None,
            "max": max(numbers) if numbers else None,
            "total": fsum(numbers) if numbers else None,
            "average": mean(numbers) if numbers else None,
            "non_finite": sum(not isfinite(value) for value in numbers),
            "non_numeric": sum(not is_missing(value) and (not isinstance(value, (int, float)) or isinstance(value, bool)) for value in columns[name]),
        }

    quantities = columns["Quantity"]
    invalid_quantities = [
        value for value in quantities
        if not is_missing(value) and (not isinstance(value, (int, float)) or isinstance(value, bool) or value <= 0 or int(value) != value)
    ]
    sales_values = columns["Sales"]
    nonpositive_sales = [value for value in sales_values if isinstance(value, (int, float)) and value <= 0]
    discounts = columns["Discount"]
    invalid_discounts = [value for value in discounts if isinstance(value, (int, float)) and not 0 <= value <= 1]
    profits = columns["Profit"]
    negative_profits = [value for value in profits if isinstance(value, (int, float)) and value < 0]
    row_id_counts = Counter(value for value in columns["Row ID"] if not is_missing(value))
    duplicate_row_ids = sum(count > 1 for count in row_id_counts.values())
    region_values = {value for value in columns["Region"] if not is_missing(value)}
    text_whitespace_values = [
        (name, value)
        for name in headers
        for value in columns[name]
        if isinstance(value, str) and value != value.strip()
    ]
    whitespace_details: dict[str, Counter[str]] = defaultdict(Counter)
    for name, value in text_whitespace_values:
        whitespace_details[name][value] += 1

    profile_lines = [
        "# Dataset Profile",
        "",
        "## 1. Source Dataset",
        f"- **File:** `{workbook_path.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- **Sheet:** `{SHEET_NAME}` (only sheet: {', '.join(f'`{name}`' for name in workbook.sheetnames)})",
        f"- **Rows:** {len(rows):,} data rows, excluding the header",
        f"- **Columns:** {len(headers)}",
        "- **Workbook access:** Opened successfully in read-only mode; the source file was not written.",
        "",
        "## 2. Column Dictionary",
        "",
        "| Column name | Observed data type | Description/purpose | Missing values | Unique values |",
        "|---|---|---|---:|---:|",
    ]
    for name in headers:
        profile_lines.append(
            f"| {name} | {infer_type(columns[name])} | {DESCRIPTIONS[name]} | "
            f"{missing_counts[name]:,} | {unique_counts[name]:,} |"
        )

    profile_lines.extend([
        "",
        "## 3. Data Quality",
        "",
        f"- **Missing values:** {sum(missing_counts.values()):,} missing cells across all columns. Per-column counts are shown above.",
        f"- **Completely duplicate data rows:** {duplicate_rows:,} duplicate occurrences (rows identical across all 21 fields).",
        f"- **Duplicate `Row ID` values:** {duplicate_row_ids:,} duplicated identifier values; `Row ID` is expected to identify a source row.",
        f"- **Repeated `Order ID` values:** {orders_with_multiple_rows:,} of {order_count:,} orders appear on multiple rows; this is consistent with orders containing multiple product lines, not a uniqueness failure.",
        f"- **Repeated `Customer ID` values:** {customers_with_multiple_rows:,} customers appear on multiple rows; {customers_with_multiple_orders:,} customers occur in more than one distinct order.",
        f"- **Repeated `Product ID` values:** {products_with_multiple_rows:,} products appear on multiple rows; {products_with_multiple_orders:,} products occur in more than one distinct order.",
        f"- **Invalid/missing dates:** Order Date has {invalid_order_dates:,} non-date values; Ship Date has {invalid_ship_dates:,} non-date values.",
        f"- **Ship date before order date:** {len(ship_before_order):,} rows.",
        f"- **Quantity checks:** {len(invalid_quantities):,} non-missing values are non-positive, non-numeric, or non-integer.",
        f"- **Sales checks:** {len(nonpositive_sales):,} rows have zero or negative sales; {numeric_stats['Sales']['non_numeric']:,} non-numeric and {numeric_stats['Sales']['non_finite']:,} non-finite values.",
        f"- **Discount checks:** {len(invalid_discounts):,} numeric values fall outside the expected 0 to 1 interval; {numeric_stats['Discount']['non_numeric']:,} values are non-numeric; observed values are {format_values(set(discounts))}.",
        f"- **Profit checks:** {len(negative_profits):,} rows have negative profit (losses; these are analytically meaningful, not automatically invalid); {numeric_stats['Profit']['non_numeric']:,} values are non-numeric.",
        f"- **Text formatting:** {len(text_whitespace_values):,} text values have leading or trailing whitespace.",
        f"- **Region values:** {format_values(region_values)} ({'constant' if len(region_values) <= 1 else 'not constant'} across non-missing rows).",
        f"- **Customer ID to Customer Name consistency:** {len(customer_name_conflicts):,} IDs map to multiple names.",
        f"- **Customer ID to Segment consistency:** {len(customer_segment_conflicts):,} IDs map to multiple segments.",
        f"- **Product ID to Product Name consistency:** {len(product_name_conflicts):,} IDs map to multiple names.",
        f"- **Product ID to Category consistency:** {len(product_category_conflicts):,} IDs map to multiple categories.",
        f"- **Product ID to Sub-Category consistency:** {len(product_subcategory_conflicts):,} IDs map to multiple sub-categories.",
        f"- **Within-order consistency:** {len(inconsistent_orders):,} orders vary in customer, order date, ship date, or ship mode across their rows.",
        "- **Categorical domain review:** observed values are listed in Dataset Statistics. The workbook and plan do not provide an authoritative allowed-value list, so values are reported rather than labeled unexpected by assumption.",
        "",
        "## 4. Dataset Statistics",
        "",
        "| Measure | Result |",
        "|---|---:|",
        f"| Unique customers | {unique_counts['Customer ID']:,} |",
        f"| Unique products | {unique_counts['Product ID']:,} |",
        f"| Unique orders | {unique_counts['Order ID']:,} |",
        f"| States | {unique_counts['State']:,} |",
        f"| Cities | {unique_counts['City']:,} |",
        f"| Shipping modes | {unique_counts['Ship Mode']:,} |",
        f"| Customer segments | {unique_counts['Segment']:,} |",
        f"| Product categories | {unique_counts['Category']:,} |",
        f"| Product sub-categories | {unique_counts['Sub-Category']:,} |",
        f"| Order Date range | {min(order_dates).date().isoformat() if order_dates else 'n/a'} to {max(order_dates).date().isoformat() if order_dates else 'n/a'} |",
        f"| Ship Date range | {min(ship_dates).date().isoformat() if ship_dates else 'n/a'} to {max(ship_dates).date().isoformat() if ship_dates else 'n/a'} |",
        f"| Total Sales | {markdown_number(numeric_stats['Sales']['total'])} |",
        f"| Total Quantity | {markdown_number(numeric_stats['Quantity']['total'])} |",
        f"| Total Profit | {markdown_number(numeric_stats['Profit']['total'])} |",
        f"| Average Sales per row | {markdown_number(numeric_stats['Sales']['average'])} |",
        f"| Average Profit per row | {markdown_number(numeric_stats['Profit']['average'])} |",
        f"| Average Discount per row | {markdown_number(numeric_stats['Discount']['average'])} |",
        "",
        "Numeric ranges:",
        "",
        "| Column | Minimum | Maximum |",
        "|---|---:|---:|",
    ])
    for name in ("Sales", "Quantity", "Discount", "Profit"):
        profile_lines.append(
            f"| {name} | {markdown_number(numeric_stats[name]['min'])} | {markdown_number(numeric_stats[name]['max'])} |"
        )
    profile_lines.extend(["", "Observed categorical values:", "", "| Column | Values |", "|---|---|"])
    for name in ("Country", "State", "Ship Mode", "Segment", "Category", "Sub-Category", "Region"):
        profile_lines.append(f"| {name} | {format_values(set(columns[name]), limit=30)} |")

    profile_lines.extend([
        "",
        "## 5. Relationship and Consistency Checks",
        "",
        f"- An order can have multiple rows: **yes**; {orders_with_multiple_rows:,} orders have multiple line rows. The maximum number of rows for one order is {max(Counter(row[order_index] for row in rows).values(), default=0):,}.",
        f"- Customers can appear across multiple orders: **yes**; {customers_with_multiple_orders:,} customer IDs appear in more than one order.",
        f"- Products can appear across multiple orders: **yes**; {products_with_multiple_orders:,} product IDs appear in more than one order.",
        f"- Ship Date is on or after Order Date: **{'yes for all comparable rows' if not ship_before_order else 'no'}**; violations: {len(ship_before_order):,}.",
        f"- Quantity is a positive whole number: **{'yes for all populated values' if not invalid_quantities and missing_counts['Quantity'] == 0 else 'no'}**; invalid values: {len(invalid_quantities):,}; missing: {missing_counts['Quantity']:,}.",
        f"- Region is constant: **{'yes' if len(region_values) <= 1 else 'no'}**; values: {format_values(region_values)}.",
        f"- Customer and product mapping conflict counts are listed in Data Quality; within-order conflicts: {len(inconsistent_orders):,}.",
        "",
        "### First Five Data Rows",
        "",
        f"| {' | '.join(headers)} |",
        f"| {' | '.join(['---'] * len(headers))} |",
    ])
    for row in rows[:5]:
        profile_lines.append(f"| {' | '.join(display(value) for value in row)} |")
    profile_lines.extend([
        "",
        "### Last Five Data Rows",
        "",
        f"| {' | '.join(headers)} |",
        f"| {' | '.join(['---'] * len(headers))} |",
    ])
    for row in rows[-5:]:
        profile_lines.append(f"| {' | '.join(display(value) for value in row)} |")

    if whitespace_details:
        profile_lines.extend(["", "### Text Values With Leading or Trailing Whitespace", "", "| Column | Exact value (`repr`) | Occurrences |", "|---|---|---:|"])
        for name, values in sorted(whitespace_details.items()):
            for value, count in sorted(values.items()):
                exact_value = repr(value).replace("|", "\\|")
                profile_lines.append(f"| {name} | `{exact_value}` | {count:,} |")

    if product_name_conflicts:
        profile_lines.extend(["", "### Product ID to Product Name Conflicts", "", "| Product ID | Observed names |", "|---|---|"])
        for product_id, names in sorted(product_name_conflicts.items()):
            profile_lines.append(f"| {display(product_id)} | {format_values(names, limit=30)} |")

    findings = [
        f"- The source contains {len(rows):,} order-line rows and {unique_counts['Order ID']:,} distinct orders.",
        f"- {unique_counts['Customer ID']:,} distinct customers and {unique_counts['Product ID']:,} distinct products recur across order lines.",
        f"- The workbook reports {len(region_values)} non-missing region value(s): {format_values(region_values)}.",
        f"- There are {len(negative_profits):,} negative-profit rows; losses should be retained for business analysis.",
    ]
    for label, conflicts in (
        ("Customer names", customer_name_conflicts),
        ("Customer segments", customer_segment_conflicts),
        ("Product names", product_name_conflicts),
        ("Product categories", product_category_conflicts),
        ("Product sub-categories", product_subcategory_conflicts),
    ):
        if conflicts:
            findings.append(f"- {label} have {len(conflicts):,} identifier mapping conflict(s); inspect before treating these attributes as unique dimension values.")
    mapping_conflicts = (
        customer_name_conflicts
        or customer_segment_conflicts
        or product_name_conflicts
        or product_category_conflicts
        or product_subcategory_conflicts
    )
    invalid_numeric_values = any(stats["non_numeric"] or stats["non_finite"] for stats in numeric_stats.values())
    if (
        duplicate_rows or duplicate_row_ids or ship_before_order or invalid_quantities or invalid_discounts
        or nonpositive_sales or inconsistent_orders or mapping_conflicts or text_whitespace_values
        or invalid_numeric_values
    ):
        findings.append("- One or more data-quality checks flagged records; see the detailed counts above before loading or transforming the data.")
    else:
        findings.append("- No duplicate rows, identifier mapping conflicts, temporal violations, or out-of-range checks were detected in the tested fields.")
    profile_lines.extend(["", "## 6. Important Findings", "", *findings, ""])

    assumptions = [
        "- Each populated source row is treated as one order line; `Order ID` is therefore expected to repeat.",
        "- `Row ID` is treated as a row identifier and checked for uniqueness, not as a business entity key.",
        "- Repeated customer and product IDs are expected because the same entities participate in multiple order lines.",
        "- Negative profit represents a possible loss and is retained; it is not classified as invalid solely because it is negative.",
        "- Discount rates are checked against the conventional 0 to 1 interval; the check does not reinterpret or alter source values.",
        "- No source values are changed by this profiler. Any future handling of conflicts, missing values, or invalid records must be justified and documented.",
    ]
    profile_lines.extend(["## 7. Assumptions for Future Phases", "", *assumptions, ""])

    report_path = PROJECT_ROOT / "docs" / "dataset_profile.md"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(profile_lines), encoding="utf-8")
    print(f"Workbook: {workbook_path}")
    print(f"Sheet: {SHEET_NAME}; rows: {len(rows)}; columns: {len(headers)}")
    print(f"Report: {report_path}")
    print(f"Missing cells: {sum(missing_counts.values())}; duplicate rows: {duplicate_rows}")
    print(f"Unique orders/customers/products: {order_count}/{customer_count}/{product_count}")
    print(f"Date order violations: {len(ship_before_order)}; invalid quantities: {len(invalid_quantities)}")
    print(f"Customer-name conflicts: {len(customer_name_conflicts)}; product mapping conflicts: "
          f"{len(product_name_conflicts) + len(product_category_conflicts) + len(product_subcategory_conflicts)}")


if __name__ == "__main__":
    main()